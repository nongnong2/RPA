#차트 종류는 사이트 참고 
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

wb = load_workbook("sample4.xlsx")
ws = wb.active
bar_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)# B1:C11까지 데이터 차트 생성
bar_chart = BarChart() #차트 종류 설정(Bar, Line, Pie)
bar_chart.add_data(bar_value, titles_from_data=True) # 차트 데이터 추가, 계열에 제목 설정
bar_chart.title = "성적표"
bar_chart.style = 20 #미리 정의된 스타일 
bar_chart.y_axis.title = "점수" #y축 제목
bar_chart.x_axis.title = "번호" #x축 제목

ws.add_chart(bar_chart, "E1") # 차트 넣을 위치 정의

wb.save("sample_chart.xlsx")
